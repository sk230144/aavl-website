#!/usr/bin/env python3
"""
AAVL 99 - beta calculator.
Downloads 1 year of daily prices (25-Sep-2025 to 25-Sep-2026) for the 99 stocks and the
Nifty 50 (^NSEI) from Yahoo Finance, computes each stock's beta, and writes it into
AAVL99_Portfolio.xlsx (sheet 'Portfolio', columns M and N).

    beta = covariance(stock daily return, Nifty daily return) / variance(Nifty daily return)

Stocks with fewer than 60 common trading days (recent listings) are left blank.
"""
import os, sys, datetime as dt

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from aavl99_terminal import load_constituents   # noqa: E402

BENCH = "^NSEI"
START, END = dt.date(2025, 9, 25), dt.date(2026, 9, 25)
MIN_DAYS = 60
XLSX = os.path.join(HERE, "AAVL99_Portfolio.xlsx")


def download(tickers):
    import yfinance as yf
    import pandas as pd
    df = yf.download(sorted(set(tickers)), start=START.isoformat(),
                     end=(END + dt.timedelta(days=1)).isoformat(), interval="1d",
                     auto_adjust=True, progress=False, threads=True,
                     group_by="column", multi_level_index=True)
    if df is None or len(df) == 0:
        return pd.DataFrame()
    close = df["Close"] if isinstance(df.columns, pd.MultiIndex) else df[["Close"]]
    if close.shape[1] == 1 and len(set(tickers)) == 1:
        close.columns = list(set(tickers))
    close.index = pd.to_datetime(close.index).date
    return close


def compute_betas(cons, prices_for):
    """prices_for(list_of_tickers) -> DataFrame of adjusted closes (columns = tickers)."""
    import pandas as pd
    bench = prices_for([BENCH])
    if BENCH not in bench.columns or bench[BENCH].notna().sum() < MIN_DAYS:
        raise RuntimeError("Could not download Nifty 50 (^NSEI) prices. Check your internet and try again.")
    pending = {c["no"]: [s + ".NS" for s in c["symbols"]] + [c["symbol"] + ".BO"] for c in cons}
    found, frames = {}, [bench]
    for _ in range(4):
        batch = {no: cands.pop(0) for no, cands in pending.items() if cands}
        if not batch:
            break
        px = prices_for(list(batch.values()))
        for no, t in batch.items():
            if t in px.columns and px[t].notna().sum() > 0:
                found[no] = t
                pending.pop(no, None)
        frames.append(px[[t for t in px.columns if t in found.values()]])
        if not pending:
            break
    allpx = pd.concat(frames, axis=1)
    allpx = allpx.loc[:, ~allpx.columns.duplicated()]
    out = {}
    for c in cons:
        t = found.get(c["no"])
        if t is None:
            out[c["no"]] = (None, 0, "ticker not found")
            continue
        pair = allpx[[t, BENCH]].dropna()
        rets = pair.pct_change().dropna()
        n = len(rets)
        if n < MIN_DAYS:
            out[c["no"]] = (None, n, f"only {n} days of history")
            continue
        var = rets[BENCH].var()
        beta = rets[t].cov(rets[BENCH]) / var if var else None
        out[c["no"]] = (beta, n, "")
    return out


def write_excel(cons, betas, path=XLSX):
    from openpyxl import load_workbook
    from openpyxl.styles import Font
    wb = load_workbook(path)
    ws = wb["Portfolio"]
    row_of = {}
    for r in range(6, 6 + 150):
        v = ws.cell(r, 3).value
        if v is None:
            break
        row_of[str(v).strip().upper()] = r
    written = 0
    for c in cons:
        r = row_of.get(c["symbol"].upper())
        if r is None:
            continue
        beta, n, _ = betas[c["no"]]
        ws.cell(r, 13).value = round(beta, 4) if beta is not None else None
        ws.cell(r, 14).value = n
        ws.cell(r, 13).font = Font(name="Arial", color="0000FF")
        written += beta is not None
    s = wb["Summary"]
    s["B14"] = (f"Beta calculated on {dt.date.today():%d-%b-%Y}: daily returns {START:%d-%b-%Y} to "
                f"{END:%d-%b-%Y} vs Nifty 50, source Yahoo Finance.")
    s["B14"].font = Font(name="Arial", italic=True, color="595959")
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return written


def main():
    if not os.path.exists(XLSX):
        print(f"AAVL99_Portfolio.xlsx not found. Put it in this folder: {HERE}"); sys.exit(1)
    try:
        import yfinance  # noqa: F401
    except ImportError:
        print("yfinance is not installed. Run:  pip install yfinance pandas openpyxl"); sys.exit(1)
    cons = load_constituents()
    print("Downloading 1 year of prices for 99 stocks and Nifty 50... (about a minute)")
    betas = compute_betas(cons, download)
    try:
        n = write_excel(cons, betas)
    except PermissionError:
        print("AAVL99_Portfolio.xlsx is open in Excel. Close it and run this again."); sys.exit(1)
    shares = {c["no"]: c["shares"] * c["src_price"] for c in cons}
    tot = sum(shares.values())
    have = [(shares[no] / tot, b[0]) for no, b in betas.items() if b[0] is not None]
    cover = sum(w for w, _ in have)
    pb = sum(w * b for w, b in have) / cover if cover else float("nan")
    print(f"\nBeta found for {n} of {len(cons)} stocks ({cover*100:.1f}% of portfolio weight).")
    print(f"Portfolio beta (approx., index weights): {pb:.2f}   <- exact value is on the Summary sheet")
    missing = [(c["symbol"], betas[c['no']][2]) for c in cons if betas[c["no"]][0] is None]
    if missing:
        print("No beta for: " + ", ".join(f"{s} ({why})" for s, why in missing))
    print(f"\nSaved into {XLSX}. Open it and see Summary and Sectors.")


if __name__ == "__main__":
    main()
