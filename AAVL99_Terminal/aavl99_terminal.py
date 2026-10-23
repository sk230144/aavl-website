#!/usr/bin/env python3
"""
AAVL 99 Terminal - live market-cap weighted index of 99 NSE stocks.
Base value 1000 on 25-Sep-2026.

    python aavl99_terminal.py               -> live dashboard at http://127.0.0.1:8099
    python aavl99_terminal.py --demo        -> same dashboard with simulated prices (no internet needed)
    python aavl99_terminal.py --update-excel-> write daily closes into AAVL_99.xlsx and exit

Index formula:  Index = 1000 x SUM(shares_i x price_i) / SUM(shares_i x base_close_i)
shares_i = source_mcap_cr / source_price  (from constituents.csv, i.e. your Portfolio Data sheet)
"""
import csv, json, math, os, sys, threading, time, traceback, webbrowser, random, argparse
import datetime as dt
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from zoneinfo import ZoneInfo

HERE = os.path.dirname(os.path.abspath(__file__))
IST = ZoneInfo("Asia/Kolkata")
BASE_DATE = dt.date(2026, 9, 25)
BASE_VALUE = 1000.0
PORT = 8099
LIVE_REFRESH_SEC = 60          # during market hours
IDLE_REFRESH_SEC = 900         # outside market hours
MISMATCH_LOW, MISMATCH_HIGH = 0.75, 1.33


# ---------------------------------------------------------------- constituents
def load_constituents(path=os.path.join(HERE, "constituents.csv")):
    rows = []
    with open(path, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            mcap, price = float(r["source_mcap_cr"]), float(r["source_price"])
            if price <= 0 or mcap <= 0:
                raise ValueError(f"Row {r['no']} ({r['company']}): mcap and price must be > 0")
            syms = [s.strip() for s in r["nse_symbols"].split("|") if s.strip()]
            rows.append(dict(no=int(r["no"]), company=r["company"].strip(), symbols=syms,
                             symbol=syms[0], shares=mcap / price, src_price=price))
    if not rows:
        raise ValueError("constituents.csv is empty")
    return rows


def now_ist():
    return dt.datetime.now(IST)


def market_open_now(now=None):
    now = now or now_ist()
    if now.weekday() >= 5:
        return False
    t = now.time()
    return dt.time(9, 15) <= t <= dt.time(15, 30)


# ---------------------------------------------------------------- data sources
class YahooSource:
    """Prices from Yahoo Finance via yfinance (NSE '.NS', BSE '.BO' fallback)."""
    name = "Yahoo Finance (NSE)"

    def __init__(self, cons):
        import yfinance  # noqa: F401  (fail early with a clear message)
        self.cons = cons
        self.ticker = {}      # constituent no -> working yahoo ticker

    @staticmethod
    def _field(df, field):
        import pandas as pd
        if df is None or len(df) == 0:
            return pd.DataFrame()
        if isinstance(df.columns, pd.MultiIndex):
            lv0 = df.columns.get_level_values(0)
            if field in lv0:
                out = df[field]
            else:                                   # group_by='ticker' layout
                out = df.xs(field, axis=1, level=1)
        else:
            out = df[[field]] if field in df.columns else pd.DataFrame()
        return out

    def _download(self, tickers, **kw):
        import yfinance as yf
        return yf.download(tickers, auto_adjust=False, progress=False, threads=True,
                           group_by="column", multi_level_index=True, **kw)

    def resolve(self, start, end):
        """Find a working ticker for every constituent; return daily Close frame."""
        import pandas as pd
        pending = {c["no"]: [s + ".NS" for s in c["symbols"]] + [c["symbol"] + ".BO"] for c in self.cons}
        frames = []
        for _ in range(4):                          # try 1st candidate, then fallbacks
            batch = {no: cands.pop(0) for no, cands in pending.items() if cands}
            if not batch:
                break
            tick = sorted(set(batch.values()))
            df = self._download(tick, start=start, end=end, interval="1d")
            close = self._field(df, "Close")
            if len(tick) == 1 and close.shape[1] == 1:
                close.columns = tick
            got = [t for t in close.columns if close[t].notna().any()] if len(close) else []
            for no, t in batch.items():
                if t in got:
                    self.ticker[no] = t
                    pending.pop(no, None)
            if got:
                frames.append(close[got])
            if not pending:
                break
        daily = pd.concat(frames, axis=1) if frames else pd.DataFrame()
        daily = daily.loc[:, ~daily.columns.duplicated()]
        daily.index = pd.to_datetime(daily.index).date
        return daily

    def intraday(self):
        import pandas as pd
        tick = sorted(set(self.ticker.values()))
        if not tick:
            return pd.DataFrame()
        df = self._download(tick, period="1d", interval="1m", ignore_tz=False)
        close = self._field(df, "Close")
        if len(tick) == 1 and close.shape[1] == 1:
            close.columns = tick
        if len(close):
            idx = pd.to_datetime(close.index)
            idx = idx.tz_localize("UTC") if idx.tz is None else idx
            close.index = idx.tz_convert(IST)
        return close


class DemoSource:
    """Simulated prices so the dashboard can be tried without internet."""
    name = "DEMO (simulated prices)"

    def __init__(self, cons):
        self.cons = cons
        self.ticker = {c["no"]: c["symbol"] + ".NS" for c in cons}
        self.rng = random.Random(7)
        self.vol = {t: self.rng.uniform(0.008, 0.025) for t in self.ticker.values()}

    def resolve(self, start, end):
        import pandas as pd
        days, d = [], BASE_DATE - dt.timedelta(days=3)
        today = now_ist().date()
        while d <= today:
            if d.weekday() < 5:
                days.append(d)
            d += dt.timedelta(days=1)
        data = {}
        for c in self.cons:
            t = self.ticker[c["no"]]; p = c["src_price"]; col = []
            for day in days:
                if day > BASE_DATE:
                    p *= math.exp(self.rng.gauss(0.0004, self.vol[t]))
                col.append(round(p, 2))
            data[t] = col
        return pd.DataFrame(data, index=days)

    def intraday(self):
        import pandas as pd
        now = now_ist()
        day = now.date() if now.weekday() < 5 else now.date() - dt.timedelta(days=now.weekday() - 4)
        start = dt.datetime.combine(day, dt.time(9, 15), IST)
        end = min(now, dt.datetime.combine(day, dt.time(15, 30), IST)) if day == now.date() else \
            dt.datetime.combine(day, dt.time(15, 30), IST)
        if end <= start:
            return pd.DataFrame()
        idx = pd.date_range(start, end, freq="1min")
        prev = self._prev
        data = {}
        r = random.Random(day.toordinal())
        for t, p0 in prev.items():
            steps = [r.gauss(0, self.vol[t] / 19) for _ in idx]
            p, col = p0, []
            for s in steps:
                p *= math.exp(s); col.append(round(p, 2))
            data[t] = col
        return pd.DataFrame(data, index=idx)


# ---------------------------------------------------------------- engine
class Engine:
    def __init__(self, source, cons):
        self.src, self.cons = source, cons
        self.lock = threading.Lock()
        self.state = {"status": "loading", "message": "Fetching prices…", "source": source.name}
        self.daily = None
        self.last_daily_fetch = 0.0

    # ---------- helpers
    def _fetch_daily(self):
        start = BASE_DATE - dt.timedelta(days=10)
        end = now_ist().date() + dt.timedelta(days=1)
        daily = self.src.resolve(start.isoformat(), end.isoformat())
        if isinstance(self.src, DemoSource):
            self.src._prev = {t: daily[t].iloc[-2] if len(daily) > 1 else daily[t].iloc[-1]
                              for t in daily.columns}
        self.daily = daily
        self.last_daily_fetch = time.time()

    def compute(self):
        import pandas as pd
        if self.daily is None or time.time() - self.last_daily_fetch > (600 if market_open_now() else 3600):
            self._fetch_daily()
        daily = self.daily
        cons, tick = self.cons, self.src.ticker
        warnings = []

        # base prices
        base, shares = {}, {}
        for c in cons:
            no, t = c["no"], tick.get(c["no"])
            shares[no] = c["shares"]
            bp = None
            if t is not None and t in daily.columns and BASE_DATE in daily.index:
                v = daily.at[BASE_DATE, t]
                if pd.notna(v):
                    bp = float(v)
            if bp is None:
                bp = c["src_price"]
                warnings.append(f"{c['symbol']}: no NSE close for {BASE_DATE:%d-%b-%Y}; using sheet price ₹{bp:,.2f} as base.")
            base[no] = bp
            if t is None:
                warnings.append(f"{c['symbol']} ({c['company']}): ticker not found on Yahoo – held flat at base price. Fix nse_symbols in constituents.csv.")
            ratio = bp / c["src_price"]
            if not (MISMATCH_LOW <= ratio <= MISMATCH_HIGH):
                warnings.append(f"{c['symbol']}: 25-Sep close ₹{bp:,.2f} vs sheet price ₹{c['src_price']:,.2f} "
                                f"({(ratio-1)*100:+.0f}%). Check the sheet's market cap/price for this stock.")
        base_mcap = sum(shares[n] * base[n] for n in base)

        # daily closes (forward-filled), from base date on
        cols = {c["no"]: tick.get(c["no"]) for c in cons}
        dd = daily[daily.index >= BASE_DATE].sort_index() if len(daily) else daily
        dd = dd.ffill() if len(dd) else dd

        def price_row(row):
            out = {}
            for no, t in cols.items():
                v = row.get(t) if t is not None else None
                out[no] = float(v) if v is not None and pd.notna(v) else base[no]
            return out

        daily_points = []
        for d, row in dd.iterrows():
            px = price_row(row)
            val = BASE_VALUE * sum(shares[n] * px[n] for n in px) / base_mcap
            daily_points.append({"time": d.isoformat(), "value": round(val, 2), "_px": px})
        if not daily_points:
            daily_points.append({"time": BASE_DATE.isoformat(), "value": BASE_VALUE, "_px": dict(base)})
        official_closes = {p["time"]: p["_px"] for p in daily_points}   # exchange daily closes only

        # intraday
        intra = self.src.intraday()
        session_date = None
        intra_points = []
        if len(intra):
            session_date = intra.index[-1].date()
            prev_rows = [p for p in daily_points if p["time"] < session_date.isoformat()]
            prev_px = prev_rows[-1]["_px"] if prev_rows else dict(base)
            intra = intra[[t for t in intra.columns]].ffill()
            for ts, row in intra.iterrows():
                px = {}
                for no, t in cols.items():
                    v = row.get(t) if t is not None else None
                    px[no] = float(v) if v is not None and pd.notna(v) else prev_px[no]
                val = BASE_VALUE * sum(shares[n] * px[n] for n in px) / base_mcap
                intra_points.append({"time": int(ts.timestamp()) + 19800, "value": round(val, 2), "_px": px})

        # latest snapshot
        if intra_points:
            last_px = intra_points[-1]["_px"]; last_val = intra_points[-1]["value"]
            prev_rows = [p for p in daily_points if p["time"] < session_date.isoformat()]
            # make the daily series end with the live value for the session day
            daily_points = [p for p in daily_points if p["time"] < session_date.isoformat()]
            daily_points.append({"time": session_date.isoformat(), "value": last_val, "_px": last_px})
        else:
            last_px = daily_points[-1]["_px"]; last_val = daily_points[-1]["value"]
            prev_rows = daily_points[:-1]
            session_date = dt.date.fromisoformat(daily_points[-1]["time"])
        prev_px = prev_rows[-1]["_px"] if prev_rows else dict(base)
        prev_val = prev_rows[-1]["value"] if prev_rows else BASE_VALUE

        cur_mcap = sum(shares[n] * last_px[n] for n in last_px)
        stocks = []
        for c in cons:
            n = c["no"]; p = last_px[n]; pp = prev_px[n]
            stocks.append({
                "no": n, "symbol": c["symbol"], "company": c["company"], "ticker": tick.get(n),
                "price": round(p, 2), "prev": round(pp, 2),
                "chg_pct": round((p / pp - 1) * 100, 2) if pp else 0.0,
                "since_base_pct": round((p / base[n] - 1) * 100, 2),
                "weight_pct": round(shares[n] * p / cur_mcap * 100, 3),
                "contrib_pts": round(BASE_VALUE * shares[n] * (p - pp) / base_mcap, 2),
                "mcap_cr": round(shares[n] * p, 0),
            })
        adv = sum(1 for s in stocks if s["chg_pct"] > 0); dec = sum(1 for s in stocks if s["chg_pct"] < 0)
        now = now_ist()
        state = {
            "status": "ok", "source": self.src.name, "demo": isinstance(self.src, DemoSource),
            "updated": now.strftime("%d-%b-%Y %H:%M:%S IST"),
            "market_open": market_open_now(now),
            "session_date": session_date.isoformat() if session_date else None,
            "value": round(last_val, 2), "prev_close": round(prev_val, 2),
            "chg": round(last_val - prev_val, 2),
            "chg_pct": round((last_val / prev_val - 1) * 100, 2) if prev_val else 0.0,
            "since_base_pct": round((last_val / BASE_VALUE - 1) * 100, 2),
            "day_high": max((p["value"] for p in intra_points), default=last_val),
            "day_low": min((p["value"] for p in intra_points), default=last_val),
            "base_date": BASE_DATE.isoformat(), "base_value": BASE_VALUE,
            "base_mcap_cr": round(base_mcap, 0), "mcap_cr": round(cur_mcap, 0),
            "advances": adv, "declines": dec, "unchanged": len(stocks) - adv - dec,
            "intraday": [{"time": p["time"], "value": p["value"]} for p in intra_points],
            "daily": [{"time": p["time"], "value": p["value"]} for p in daily_points],
            "stocks": stocks, "warnings": warnings,
            "tickers_found": sum(1 for c in cons if tick.get(c["no"])), "tickers_total": len(cons),
        }
        self._closes = official_closes
        self._base = base
        with self.lock:
            self.state = state
        return state

    def refresh_safe(self):
        try:
            self.compute()
        except Exception as e:
            traceback.print_exc()
            with self.lock:
                self.state = dict(self.state, status="error" if self.state.get("status") != "ok" else "ok",
                                  error=f"{type(e).__name__}: {e}")

    def loop(self):
        while True:
            self.refresh_safe()
            time.sleep(LIVE_REFRESH_SEC if market_open_now() else IDLE_REFRESH_SEC)

    # ---------- Excel
    def update_excel(self, path=os.path.join(HERE, "AAVL_99.xlsx")):
        from openpyxl import load_workbook
        if not os.path.exists(path):
            return False, f"AAVL_99.xlsx not found next to this program ({path})."
        self.last_daily_fetch = 0          # force fresh official closes
        self.compute()
        now = now_ist()
        wb = load_workbook(path)
        if "AAVL 99" not in wb.sheetnames:
            return False, "Sheet 'AAVL 99' not found in the workbook."
        ws = wb["AAVL 99"]
        date_col = {}
        for col in range(7, ws.max_column + 1):
            v = ws.cell(11, col).value
            if hasattr(v, "year") and hasattr(v, "month") and hasattr(v, "day"):
                date_col[dt.date(v.year, v.month, v.day)] = col
        sym_row = {}
        for r in range(12, 12 + 200):
            s = ws.cell(r, 3).value
            if s is None:
                break
            sym_row[str(s).strip().upper()] = r
        written, skipped = [], []
        closes = dict(self._closes)
        closes[BASE_DATE.isoformat()] = self._base
        for dstr, px in sorted(closes.items()):
            d = dt.date.fromisoformat(dstr)
            if d == now.date() and now.time() < dt.time(16, 0):
                skipped.append(f"{d:%d-%b} (official close available after 4 PM)")
                continue
            if d not in date_col:
                skipped.append(f"{d:%d-%b} (no column)")
                continue
            col = date_col[d]
            for c in self.cons:
                r = sym_row.get(c["symbol"].upper())
                if r is None:
                    continue
                ws.cell(r, col).value = round(px[c["no"]], 2)
            written.append(f"{d:%d-%b}")
        wb.calculation.fullCalcOnLoad = True
        try:
            wb.save(path)
        except PermissionError:
            return False, "AAVL_99.xlsx is open in Excel. Close it and click again."
        msg = "Saved closes for: " + (", ".join(written) if written else "none")
        if skipped:
            msg += ". Skipped: " + ", ".join(skipped)
        return True, msg


# ---------------------------------------------------------------- web server
def lan_ip():
    """This computer's address on the Wi-Fi/LAN (no data is sent)."""
    import socket
    try:
        sk = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        sk.connect(("10.255.255.255", 1))
        ip = sk.getsockname()[0]
        sk.close()
        return ip
    except OSError:
        return None


OWNER_IPS = set()     # this computer's own addresses; filled in main()


def is_owner_ip(ip):
    return ip.startswith("127.") or ip in ("::1", "localhost") or ip in OWNER_IPS


def make_handler(engine, share_url=None):
    web = os.path.join(HERE, "web")
    files = {"/": ("index.html", "text/html; charset=utf-8"),
             "/index.html": ("index.html", "text/html; charset=utf-8"),
             "/lwc.js": ("lightweight-charts.standalone.production.js", "application/javascript")}

    class H(BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def _send(self, code, body, ctype):
            self.send_response(code)
            self.send_header("Content-Type", ctype)
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            path = self.path.split("?")[0]
            if path == "/api/state":
                owner = is_owner_ip(self.client_address[0])
                with engine.lock:
                    st = dict(engine.state, owner=owner, share_url=share_url if owner else None)
                return self._send(200, json.dumps(st).encode(), "application/json")
            if path in files:
                fn, ct = files[path]
                with open(os.path.join(web, fn), "rb") as f:
                    return self._send(200, f.read(), ct)
            self._send(404, b"Not found", "text/plain")

        def do_POST(self):
            path = self.path.split("?")[0]
            if not is_owner_ip(self.client_address[0]):
                return self._send(403, b'{"ok":false,"message":"View only"}', "application/json")
            if path == "/api/refresh":
                engine.refresh_safe()
                return self._send(200, b'{"ok":true}', "application/json")
            if path == "/api/save-excel":
                try:
                    ok, msg = engine.update_excel()
                except Exception as e:
                    ok, msg = False, f"{type(e).__name__}: {e}"
                return self._send(200, json.dumps({"ok": ok, "message": msg}).encode(), "application/json")
            self._send(404, b"Not found", "text/plain")
    return H


def main():
    ap = argparse.ArgumentParser(description="AAVL 99 live index terminal")
    ap.add_argument("--demo", action="store_true", help="use simulated prices (no internet)")
    ap.add_argument("--update-excel", action="store_true", help="write daily closes into AAVL_99.xlsx and exit")
    ap.add_argument("--port", type=int, default=PORT)
    ap.add_argument("--no-browser", action="store_true")
    ap.add_argument("--local-only", action="store_true", help="do not share on Wi-Fi; only this computer can open it")
    a = ap.parse_args()

    cons = load_constituents()
    try:
        src = DemoSource(cons) if a.demo else YahooSource(cons)
    except ImportError:
        print("yfinance is not installed. Run:  pip install yfinance pandas openpyxl"); sys.exit(1)
    eng = Engine(src, cons)

    if a.update_excel:
        eng.compute()
        ok, msg = eng.update_excel(); print(msg); sys.exit(0 if ok else 1)

    threading.Thread(target=eng.loop, daemon=True).start()
    ip = None if a.local_only else lan_ip()
    if ip:
        OWNER_IPS.add(ip)
    share_url = f"http://{ip}:{a.port}" if ip else None
    host = "127.0.0.1" if a.local_only else "0.0.0.0"
    try:
        srv = ThreadingHTTPServer((host, a.port), make_handler(eng, share_url))
    except OSError:
        print(f"Port {a.port} is busy - AAVL 99 Terminal may already be running in another window."); sys.exit(1)
    url = f"http://127.0.0.1:{a.port}"
    print(f"AAVL 99 Terminal running.   Source: {src.name}")
    print(f"  On this computer : {url}")
    if share_url:
        print(f"  Share on Wi-Fi   : {share_url}   (view only for other devices)")
    elif not a.local_only:
        print("  Wi-Fi sharing    : not available (this computer is not connected to a network)")
    print("Keep this window open. Close it to stop.")
    if not a.no_browser:
        threading.Timer(1.5, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\nStopped.")


if __name__ == "__main__":
    main()
